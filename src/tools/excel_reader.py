"""Read Excel workbook structure and preview data without analyzing it."""

from __future__ import annotations

import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


def _error(
    code: str,
    message: str,
    *,
    file_name: str | None = None,
    sheet_names: list[str] | None = None,
) -> dict[str, Any]:
    result: dict[str, Any] = {
        "success": False,
        "error": {"code": code, "message": message},
    }
    if file_name is not None:
        result["file_name"] = file_name
    if sheet_names is not None:
        result["sheet_names"] = sheet_names
    return result


def _is_missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _column_name(value: Any, index: int, seen: dict[str, int]) -> str:
    base = f"unnamed_column_{index}" if _is_missing(value) else str(value).strip()
    seen[base] = seen.get(base, 0) + 1
    return base if seen[base] == 1 else f"{base}_{seen[base]}"


def _dtype(values: list[Any]) -> str:
    present = [value for value in values if not _is_missing(value)]
    if not present:
        return "unknown"
    if all(isinstance(value, bool) for value in present):
        return "boolean"
    if all(isinstance(value, datetime) for value in present):
        return "datetime"
    if all(isinstance(value, date) and not isinstance(value, datetime) for value in present):
        return "date"
    if all(isinstance(value, int) and not isinstance(value, bool) for value in present):
        return "integer"
    if all(
        isinstance(value, (int, float)) and not isinstance(value, bool)
        for value in present
    ):
        return "float"
    if all(isinstance(value, str) for value in present):
        return "string"
    return "mixed"


def _is_sensitive_column(name: str) -> bool:
    normalized = re.sub(r"[\s_-]+", "", name).lower()
    exact_names = {
        "name",
        "fullname",
        "employeename",
        "phone",
        "phonenumber",
        "mobile",
        "mobilenumber",
        "telephone",
        "email",
        "emailaddress",
        "address",
        "homeaddress",
        "idcard",
        "idcardnumber",
        "nationalid",
    }
    chinese_markers = ("姓名", "身份证", "手机号", "手机号码", "邮箱", "地址", "住址")
    return normalized in exact_names or any(marker in normalized for marker in chinese_markers)


def _serialize(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def read_excel_metadata(
    file_path: str,
    sheet_name: str | None = None,
    preview_rows: int = 5,
) -> dict[str, Any]:
    """Return workbook metadata and a privacy-safe preview, without analysis."""
    path = Path(file_path).expanduser()
    file_name = path.name

    if isinstance(preview_rows, bool) or not isinstance(preview_rows, int) or preview_rows < 0:
        return _error(
            "invalid_preview_rows",
            "preview_rows 必须是大于或等于 0 的整数。",
            file_name=file_name,
        )
    if not path.exists() or not path.is_file():
        return _error("file_not_found", f"文件不存在：{file_path}", file_name=file_name)
    if path.suffix.lower() != ".xlsx":
        return _error(
            "invalid_file_type",
            "仅支持 .xlsx 文件。",
            file_name=file_name,
        )

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException):
        return _error(
            "file_corrupted",
            "Excel 文件已损坏或不是有效的 .xlsx 文件。",
            file_name=file_name,
        )
    except Exception as exc:
        return _error(
            "parse_error",
            f"Excel 无法正常解析：{type(exc).__name__}。",
            file_name=file_name,
        )

    try:
        sheet_names = list(workbook.sheetnames)
        if not sheet_names:
            return _error(
                "empty_workbook",
                "Excel 中没有可读取的 Sheet。",
                file_name=file_name,
                sheet_names=[],
            )
        if sheet_name is None and len(sheet_names) > 1:
            return {
                "success": True,
                "file_name": file_name,
                "sheet_names": sheet_names,
                "selected_sheet": None,
                "message": "文件包含多个 Sheet，请指定 sheet_name 后再读取元数据。",
            }
        if sheet_name is not None and sheet_name not in sheet_names:
            return _error(
                "sheet_not_found",
                f"Sheet 不存在：{sheet_name}。",
                file_name=file_name,
                sheet_names=sheet_names,
            )

        selected_sheet = sheet_name or sheet_names[0]
        worksheet = workbook[selected_sheet]
        rows = list(worksheet.iter_rows(values_only=True))
        while rows and all(_is_missing(value) for value in rows[-1]):
            rows.pop()

        used_column_count = max(
            (
                index + 1
                for row in rows
                for index, value in enumerate(row)
                if not _is_missing(value)
            ),
            default=0,
        )
        if not rows or used_column_count == 0:
            return _error(
                "empty_sheet",
                f"Sheet“{selected_sheet}”为空。",
                file_name=file_name,
                sheet_names=sheet_names,
            )

        trimmed_rows = [tuple(row[:used_column_count]) for row in rows]
        seen: dict[str, int] = {}
        columns = [
            _column_name(value, index, seen)
            for index, value in enumerate(trimmed_rows[0], start=1)
        ]
        data_rows = [
            row
            for row in trimmed_rows[1:]
            if any(not _is_missing(value) for value in row)
        ]
        dtypes = {
            column: _dtype([row[index] if index < len(row) else None for row in data_rows])
            for index, column in enumerate(columns)
        }
        missing_values = {
            column: sum(
                1
                for row in data_rows
                if index >= len(row) or _is_missing(row[index])
            )
            for index, column in enumerate(columns)
        }
        sensitive_columns = [
            column for column in columns if _is_sensitive_column(column)
        ]
        preview: list[dict[str, Any]] = []
        for row in data_rows[:preview_rows]:
            preview.append(
                {
                    column: (
                        None
                        if index >= len(row) or _is_missing(row[index])
                        else "***"
                        if column in sensitive_columns
                        else _serialize(row[index])
                    )
                    for index, column in enumerate(columns)
                }
            )

        return {
            "success": True,
            "file_name": file_name,
            "sheet_names": sheet_names,
            "selected_sheet": selected_sheet,
            "row_count": len(data_rows),
            "column_count": len(columns),
            "columns": columns,
            "dtypes": dtypes,
            "missing_values": missing_values,
            "sensitive_columns_detected": sensitive_columns,
            "preview": preview,
        }
    except Exception as exc:
        return _error(
            "parse_error",
            f"Excel 无法正常解析：{type(exc).__name__}。",
            file_name=file_name,
            sheet_names=list(workbook.sheetnames),
        )
    finally:
        workbook.close()


def read_excel_data(
    file_path: str,
    sheet_name: str | None = None,
) -> dict[str, Any]:
    """Read complete worksheet records for local analytics orchestration.

    This reader performs no statistics. Its ``records`` result stays inside the
    local process; the Agent passes only aggregated analysis results to the model.
    """
    metadata = read_excel_metadata(file_path, sheet_name=sheet_name, preview_rows=0)
    if not metadata.get("success"):
        return metadata
    selected_sheet = metadata.get("selected_sheet")
    if not isinstance(selected_sheet, str):
        return _error(
            "sheet_required",
            "文件包含多个 Sheet，请先指定 sheet_name。",
            file_name=metadata.get("file_name"),
            sheet_names=metadata.get("sheet_names"),
        )

    path = Path(file_path).expanduser()
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException):
        return _error(
            "file_corrupted",
            "Excel 文件已损坏或不是有效的 .xlsx 文件。",
            file_name=path.name,
        )
    except Exception as exc:
        return _error(
            "parse_error",
            f"Excel 无法正常解析：{type(exc).__name__}。",
            file_name=path.name,
        )

    try:
        worksheet = workbook[selected_sheet]
        columns = list(metadata["columns"])
        records: list[dict[str, Any]] = []
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
            if row_index == 0:
                continue
            values = tuple(row[: len(columns)])
            if not any(not _is_missing(value) for value in values):
                continue
            records.append(
                {
                    column: values[index] if index < len(values) else None
                    for index, column in enumerate(columns)
                }
            )
        return {
            "success": True,
            "file_name": metadata["file_name"],
            "sheet_names": metadata["sheet_names"],
            "selected_sheet": selected_sheet,
            "columns": columns,
            "row_count": len(records),
            "records": records,
        }
    except Exception as exc:
        return _error(
            "parse_error",
            f"Excel 无法正常解析：{type(exc).__name__}。",
            file_name=path.name,
            sheet_names=list(workbook.sheetnames),
        )
    finally:
        workbook.close()


EXCEL_READER_TOOL: dict[str, Any] = {
    "type": "function",
    "name": "read_excel_metadata",
    "description": (
        "当用户提供本地 .xlsx 文件路径，并询问文件名、Sheet、指定 Sheet、行列数、"
        "字段名、字段类型、缺失值数量、前几行预览或敏感字段时使用。此工具只读取和"
        "描述 Excel，不进行工资、离职率、绩效、部门、异常值、相关性或任何 HR 数据分析。"
    ),
    "strict": True,
    "parameters": {
        "type": "object",
        "properties": {
            "file_path": {
                "type": "string",
                "description": "用户明确提供的本地 .xlsx 文件路径。",
            },
            "sheet_name": {
                "type": ["string", "null"],
                "description": "用户指定的 Sheet；未指定时传 null。",
            },
            "preview_rows": {
                "type": "integer",
                "minimum": 0,
                "description": "需要预览的数据行数，默认使用 5。",
            },
        },
        "required": ["file_path", "sheet_name", "preview_rows"],
        "additionalProperties": False,
    },
}
