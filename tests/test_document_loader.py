"""Tests for V0.5 document validation and standardized loading."""

from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfWriter

from src.knowledge import load_document


def _write_text_pdf(path: Path, text: str = "Quarterly performance meeting") -> None:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii")
    )
    path.write_bytes(bytes(output))


def test_loads_txt_and_returns_standard_document(tmp_path: Path) -> None:
    path = tmp_path / "policy.txt"
    path.write_text("调薪评审在七月进行。", encoding="utf-8")

    result = load_document(str(path))

    assert result["success"] is True
    assert result["document_id"].startswith("DOC-")
    assert result["file_type"] == "txt"
    assert result["title"] == "policy"
    assert result["text"] == "调薪评审在七月进行。"
    assert result["character_count"] == len(result["text"])
    assert result["parse_status"] == "ready"
    assert Path(result["source_path"]) == path.resolve()


def test_loads_markdown(tmp_path: Path) -> None:
    path = tmp_path / "handbook.md"
    path.write_text("# 员工手册\n\n只用于测试。", encoding="utf-8")

    result = load_document(str(path))

    assert result["success"] is True
    assert "员工手册" in result["text"]
    assert result["metadata"]["parser"] == "utf-8-text"


def test_loads_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "policy.docx"
    document = DocxDocument()
    document.add_heading("薪酬制度", level=1)
    document.add_paragraph("年度调薪在七月评审。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "规则"
    document.save(path)

    result = load_document(str(path))

    assert result["success"] is True
    assert result["file_type"] == "docx"
    assert "# 薪酬制度" in result["text"]
    assert "项目 | 规则" in result["text"]
    assert result["metadata"]["table_count"] == 1


def test_loads_machine_readable_pdf(tmp_path: Path) -> None:
    path = tmp_path / "policy.pdf"
    _write_text_pdf(path)

    result = load_document(str(path))

    assert result["success"] is True
    assert "Quarterly performance meeting" in result["text"]
    assert result["metadata"]["page_count"] == 1


def test_missing_and_unsupported_files_return_clear_errors(tmp_path: Path) -> None:
    missing = load_document(str(tmp_path / "missing.md"))
    unsupported_path = tmp_path / "employees.xlsx"
    unsupported_path.write_bytes(b"not relevant")
    unsupported = load_document(str(unsupported_path))

    assert missing["error"]["code"] == "file_not_found"
    assert unsupported["error"]["code"] == "unsupported_file_type"
    assert "Excel Reader" in unsupported["error"]["message"]


def test_blank_pdf_is_reported_as_scanned_or_unreadable(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as output:
        writer.write(output)

    result = load_document(str(path))

    assert result["success"] is False
    assert result["parse_status"] == "scanned_or_unreadable_pdf"
    assert "OCR" in result["error"]["message"]


def test_empty_text_file_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text(" \n", encoding="utf-8")

    result = load_document(str(path))

    assert result["error"]["code"] == "empty_document"

