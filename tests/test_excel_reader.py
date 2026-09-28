"""Tests for the V0.3 Excel metadata reader."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
import pytest

from src.tools import read_excel_data, read_excel_metadata


SAMPLE_PATH = Path(__file__).resolve().parents[1] / "data" / "sample_employees.xlsx"


def test_reads_valid_xlsx() -> None:
    result = read_excel_metadata(str(SAMPLE_PATH))

    assert result["success"] is True
    assert result["file_name"] == "sample_employees.xlsx"


def test_returns_sheet_names_and_selected_sheet() -> None:
    result = read_excel_metadata(str(SAMPLE_PATH))

    assert result["sheet_names"] == ["员工信息"]
    assert result["selected_sheet"] == "员工信息"


def test_returns_expected_columns() -> None:
    result = read_excel_metadata(str(SAMPLE_PATH))

    assert result["columns"] == [
        "employee_id",
        "department",
        "position",
        "hire_date",
        "monthly_salary",
        "performance_score",
        "employment_status",
        "termination_date",
    ]


def test_returns_row_and_column_counts() -> None:
    result = read_excel_metadata(str(SAMPLE_PATH))

    assert result["row_count"] == 20
    assert result["column_count"] == 8


def test_returns_column_data_types() -> None:
    result = read_excel_metadata(str(SAMPLE_PATH))

    assert result["dtypes"]["employee_id"] == "string"
    assert result["dtypes"]["monthly_salary"] == "float"
    assert result["dtypes"]["performance_score"] == "float"
    assert result["dtypes"]["hire_date"] == "datetime"


def test_returns_missing_value_counts() -> None:
    result = read_excel_metadata(str(SAMPLE_PATH))

    assert result["missing_values"]["monthly_salary"] == 1
    assert result["missing_values"]["performance_score"] == 1
    assert result["missing_values"]["termination_date"] == 15


def test_returns_first_five_rows_as_json_safe_preview() -> None:
    result = read_excel_metadata(str(SAMPLE_PATH))

    assert len(result["preview"]) == 5
    assert result["preview"][0]["employee_id"] == "EMP001"
    assert isinstance(result["preview"][0]["hire_date"], str)
    json.dumps(result, ensure_ascii=False)


def test_missing_file_returns_clear_error(tmp_path: Path) -> None:
    result = read_excel_metadata(str(tmp_path / "missing.xlsx"))

    assert result["success"] is False
    assert result["error"]["code"] == "file_not_found"


def test_missing_sheet_lists_available_sheets() -> None:
    result = read_excel_metadata(str(SAMPLE_PATH), sheet_name="不存在")

    assert result["success"] is False
    assert result["error"]["code"] == "sheet_not_found"
    assert result["sheet_names"] == ["员工信息"]


def test_wrong_file_type_returns_clear_error(tmp_path: Path) -> None:
    path = tmp_path / "employees.csv"
    path.write_text("employee_id\nEMP001\n", encoding="utf-8")

    result = read_excel_metadata(str(path))

    assert result["success"] is False
    assert result["error"]["code"] == "invalid_file_type"


def test_sensitive_columns_are_detected_and_preview_is_masked(tmp_path: Path) -> None:
    path = tmp_path / "sensitive.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["employee_id", "姓名", "email", "地址", "department"])
    sheet.append(
        ["EMP001", "虚构姓名", "fake@example.invalid", "虚构地址", "测试部"]
    )
    workbook.save(path)

    result = read_excel_metadata(str(path))

    assert result["sensitive_columns_detected"] == ["姓名", "email", "地址"]
    assert result["preview"][0]["姓名"] == "***"
    assert result["preview"][0]["email"] == "***"
    assert result["preview"][0]["地址"] == "***"
    assert result["preview"][0]["employee_id"] == "EMP001"
    assert result["preview"][0]["department"] == "测试部"


def test_multiple_sheets_require_explicit_selection(tmp_path: Path) -> None:
    path = tmp_path / "multiple.xlsx"
    workbook = Workbook()
    workbook.active.title = "员工信息"
    workbook.create_sheet("离职记录")
    workbook.save(path)

    result = read_excel_metadata(str(path))

    assert result == {
        "success": True,
        "file_name": "multiple.xlsx",
        "sheet_names": ["员工信息", "离职记录"],
        "selected_sheet": None,
        "message": "文件包含多个 Sheet，请指定 sheet_name 后再读取元数据。",
    }


def test_corrupted_xlsx_returns_clear_error(tmp_path: Path) -> None:
    path = tmp_path / "corrupted.xlsx"
    path.write_bytes(b"not an xlsx archive")

    result = read_excel_metadata(str(path))

    assert result["success"] is False
    assert result["error"]["code"] == "file_corrupted"


def test_unexpected_parser_error_returns_clear_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def fail_to_parse(*args: object, **kwargs: object) -> None:
        raise RuntimeError("模拟解析失败")

    monkeypatch.setattr("src.tools.excel_reader.load_workbook", fail_to_parse)

    result = read_excel_metadata(str(SAMPLE_PATH))

    assert result["success"] is False
    assert result["error"]["code"] == "parse_error"


def test_empty_sheet_returns_clear_error(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    Workbook().save(path)

    result = read_excel_metadata(str(path))

    assert result["success"] is False
    assert result["error"]["code"] == "empty_sheet"


def test_invalid_preview_rows_returns_clear_error() -> None:
    result = read_excel_metadata(str(SAMPLE_PATH), preview_rows=-1)

    assert result["success"] is False
    assert result["error"]["code"] == "invalid_preview_rows"


def test_reads_complete_records_for_local_analytics() -> None:
    result = read_excel_data(str(SAMPLE_PATH))

    assert result["success"] is True
    assert result["row_count"] == 20
    assert len(result["records"]) == 20
    assert result["records"][0]["employee_id"] == "EMP001"
