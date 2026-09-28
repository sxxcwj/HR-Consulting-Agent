"""Focused tests for format-specific V0.5 parsing behavior."""

from pathlib import Path

from docx import Document as DocxDocument

from src.knowledge.document_parser import parse_document


def test_docx_keeps_paragraph_table_paragraph_order(tmp_path: Path) -> None:
    path = tmp_path / "ordered.docx"
    document = DocxDocument()
    document.add_paragraph("第一段")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "中间"
    table.cell(0, 1).text = "表格"
    document.add_paragraph("最后一段")
    document.save(path)

    result = parse_document(path)

    assert result["success"] is True
    assert result["text"].index("第一段") < result["text"].index("中间 | 表格")
    assert result["text"].index("中间 | 表格") < result["text"].index("最后一段")


def test_corrupted_docx_returns_parse_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not-a-docx")

    result = parse_document(path)

    assert result["success"] is False
    assert result["error"]["code"] == "docx_parse_error"


def test_invalid_utf8_text_returns_decode_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.txt"
    path.write_bytes(b"\xff\xfe\xfa")

    result = parse_document(path)

    assert result["error"]["code"] == "text_decode_error"


def test_parser_rejects_xlsx_to_preserve_reader_boundary(tmp_path: Path) -> None:
    path = tmp_path / "employees.xlsx"
    path.write_bytes(b"placeholder")

    result = parse_document(path)

    assert result["error"]["code"] == "unsupported_file_type"

